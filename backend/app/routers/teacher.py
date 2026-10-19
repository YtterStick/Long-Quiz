import math
from datetime import datetime, timezone, timedelta
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query, Response, Form
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy import select, func, and_, or_, update, delete, String
from sqlalchemy.orm import selectinload, joinedload
from typing import Optional, Sequence, Union
from app.database import get_db
from app.models import (
    User, Teacher, Student, Section, Exam, ExamAnswerKey, ExamQuestion,
    ExamSubmission, SubmissionAnswer, Role, PasswordResetRequest, ResetRequestStatus,
    SectionExam, ExamLobbyEntry, TeacherExamGenerationLog
)
from app.schemas import (
    SectionCreate, SectionUpdate, SectionOut,
    StudentCreate, StudentOut, StudentBulkCreateResponse, StudentBatchCreate, EnrollStudent,
    PasswordResetRequestOut, StudentResetPasswordResponse,
    ExamCreate, ExamUpdate, ExamOut, AnswerKeyItem,
    QuestionItem, QuestionOut, QuestionnaireUpdate, ExamVisibilityUpdate,
    ExamOnlineToggle, ExamGradesReleaseToggle,
    AssignExamToSection, SectionExamOut, SectionExamOnlineToggle, SectionExamGradesToggle,
    SubmissionOut, SubmissionOverride,
    TeacherDashboardStats, ItemAnalysis, GradeItem
)
from app.auth import hash_password
from app.deps import require_role
from app.utils import (
    generate_temp_password,
    generate_username,
    generate_student_username,
    generate_student_number,
    parse_student_csv,
    parse_student_file,
    clean_student_omr_number,
    format_full_student_id,
    STATIC_STUDENT_PREFIX,
)

router = APIRouter(prefix="/api/teacher", tags=["Teacher"])


from pydantic import BaseModel


class TeacherProfileUpdateRequest(BaseModel):
    avatar_config: Optional[str] = None
    bio: Optional[str] = None


async def _get_teacher(user: User, db: AsyncSession) -> Teacher:
    """Helper to get the Teacher profile for the current user."""
    result = await db.execute(select(Teacher).where(Teacher.user_id == user.id))
    teacher = result.scalar_one_or_none()
    if not teacher:
        raise HTTPException(status_code=404, detail="Teacher profile not found")
    return teacher


@router.get("/profile")
async def get_teacher_profile(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Retrieve full teacher profile, avatar customization config, and philosophy bio."""
    teacher = await _get_teacher(current_user, db)
    return {
        "teacher_id": teacher.id,
        "first_name": current_user.first_name,
        "last_name": current_user.last_name,
        "username": current_user.username,
        "avatar_config": teacher.avatar_config,
        "bio": teacher.bio,
    }


@router.put("/profile")
async def update_teacher_profile(
    payload: TeacherProfileUpdateRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Update teacher avatar config and bio. Names/usernames are non-editable institutional records."""
    teacher = await _get_teacher(current_user, db)
    if payload.avatar_config is not None:
        teacher.avatar_config = payload.avatar_config
    if payload.bio is not None:
        teacher.bio = payload.bio.strip()[:160]
    await db.commit()
    await db.refresh(teacher)
    return {
        "ok": True,
        "avatar_config": teacher.avatar_config,
        "bio": teacher.bio,
        "first_name": current_user.first_name,
        "last_name": current_user.last_name,
    }


# ─── Dashboard ────────────────────────────────────────────────

@router.get("/dashboard", response_model=TeacherDashboardStats)
async def get_dashboard(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)

    sections = await db.execute(select(func.count(Section.id)).where(Section.teacher_id == teacher.id))
    students = await db.execute(
        select(func.count(Student.id)).join(Section).where(Section.teacher_id == teacher.id)
    )
    exams = await db.execute(select(func.count(Exam.id)).where(Exam.teacher_id == teacher.id))
    subs = await db.execute(
        select(func.count(ExamSubmission.id)).join(Exam).where(Exam.teacher_id == teacher.id)
    )

    return TeacherDashboardStats(
        total_sections=sections.scalar_one(),
        total_students=students.scalar_one(),
        total_exams=exams.scalar_one(),
        total_submissions=subs.scalar_one(),
    )


# ─── Sections ─────────────────────────────────────────────────

@router.get("/sections", response_model=list[SectionOut])
async def list_sections(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Section)
        .where(Section.teacher_id == teacher.id)
        .order_by(Section.created_at.desc())
    )
    sections = result.scalars().all()

    counts_res = await db.execute(
        select(Student.section_id, func.count(Student.id))
        .where(Student.section_id.in_([s.id for s in sections]))
        .group_by(Student.section_id)
    ) if sections else None
    counts = dict(counts_res.all()) if counts_res is not None else {}

    out = []
    for s in sections:
        out.append(SectionOut(
            id=s.id, name=s.name, subject=s.subject,
            teacher_id=s.teacher_id, created_at=s.created_at,
            student_count=counts.get(s.id, 0),
        ))
    return out


@router.post("/sections", response_model=SectionOut)
async def create_section(
    body: SectionCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    section = Section(name=body.name, subject=body.subject, teacher_id=teacher.id)
    db.add(section)
    await db.flush()
    return SectionOut(
        id=section.id, name=section.name, subject=section.subject,
        teacher_id=section.teacher_id, created_at=section.created_at,
        student_count=0,
    )


@router.put("/sections/{section_id}", response_model=SectionOut)
async def update_section(
    section_id: str,
    body: SectionUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    section = result.scalar_one_or_none()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    if body.name is not None:
        section.name = body.name
    if body.subject is not None:
        section.subject = body.subject
    await db.flush()

    count_result = await db.execute(
        select(func.count(Student.id)).where(Student.section_id == section.id)
    )
    return SectionOut(
        id=section.id, name=section.name, subject=section.subject,
        teacher_id=section.teacher_id, created_at=section.created_at,
        student_count=count_result.scalar_one(),
    )


@router.delete("/sections/{section_id}")
async def delete_section(
    section_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    section = result.scalar_one_or_none()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    # Unassign all students in this section so their accounts and submissions are preserved
    await db.execute(
        update(Student).where(Student.section_id == section_id).values(section_id=None)
    )
    await db.delete(section)
    await db.flush()
    return {"message": "Section deleted successfully"}


# ─── Section Exams (Class Exam Administration) ─────────────────

@router.get("/sections/{section_id}/exams", response_model=list[SectionExamOut])
async def list_section_exams(
    section_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """List all exams assigned to this section with online active, grade release status, and class score stats."""
    teacher = await _get_teacher(current_user, db)
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec_res.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(SectionExam)
        .options(
            selectinload(SectionExam.exam).selectinload(Exam.answer_keys),
            selectinload(SectionExam.exam).selectinload(Exam.questions),
        )
        .where(SectionExam.section_id == section_id)
        .order_by(SectionExam.created_at.desc())
    )
    section_exams = result.scalars().all()

    # Get student IDs in this section
    # Single aggregate query for all exams (avoids N+1 round trips)
    stats_res = await db.execute(
        select(
            ExamSubmission.exam_id,
            func.count(ExamSubmission.id),
            func.avg(ExamSubmission.total_score),
        )
        .where(
            ExamSubmission.student_id.in_(
                select(Student.id).where(Student.section_id == section_id)
            )
        )
        .group_by(ExamSubmission.exam_id)
    )
    stats = {row[0]: (row[1], row[2]) for row in stats_res.all()}

    out = []
    for se in section_exams:
        if not se.exam:
            continue

        sub_count, avg_raw = stats.get(se.exam_id, (0, None))
        avg_score = round(float(avg_raw), 1) if avg_raw is not None else None

        out.append(SectionExamOut(
            id=se.id,
            section_id=se.section_id,
            exam_id=se.exam_id,
            is_online_active=se.is_online_active,
            is_grades_released=se.is_grades_released,
            is_started=se.is_started,
            release_mode=se.release_mode or "FULL",
            created_at=se.created_at,
            exam=ExamOut.model_validate(se.exam),
            submission_count=sub_count,
            average_score=avg_score,
            max_score=se.exam.total_questions,
        ))

    return out


@router.post("/sections/{section_id}/exams", response_model=SectionExamOut)
async def assign_exam_to_section(
    section_id: str,
    body: AssignExamToSection,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Assign an existing exam to this class/section."""
    teacher = await _get_teacher(current_user, db)
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec_res.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    exam_res = await db.execute(
        select(Exam).where(Exam.id == body.exam_id)
    )
    exam = exam_res.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    # Check if already assigned
    existing_res = await db.execute(
        select(SectionExam).where(
            SectionExam.section_id == section_id,
            SectionExam.exam_id == body.exam_id,
        )
    )
    existing = existing_res.scalar_one_or_none()
    if existing:
        result = await db.execute(
            select(SectionExam)
            .options(
                selectinload(SectionExam.exam).selectinload(Exam.answer_keys),
                selectinload(SectionExam.exam).selectinload(Exam.questions),
            )
            .where(SectionExam.id == existing.id)
        )
        se = result.scalar_one()
        return SectionExamOut(
            id=se.id,
            section_id=se.section_id,
            exam_id=se.exam_id,
            is_online_active=se.is_online_active,
            is_grades_released=se.is_grades_released,
            is_started=se.is_started,
            release_mode=se.release_mode or "FULL",
            created_at=se.created_at,
            exam=ExamOut.model_validate(se.exam) if se.exam else None,
            submission_count=0,
            average_score=None,
            max_score=exam.total_questions,
        )

    section_exam = SectionExam(
        section_id=section_id,
        exam_id=body.exam_id,
        is_online_active=False,
        is_grades_released=False,
    )
    db.add(section_exam)
    await db.commit()

    result = await db.execute(
        select(SectionExam)
        .options(
            selectinload(SectionExam.exam).selectinload(Exam.answer_keys),
            selectinload(SectionExam.exam).selectinload(Exam.questions),
        )
        .where(SectionExam.id == section_exam.id)
    )
    se = result.scalar_one()
    return SectionExamOut(
        id=se.id,
        section_id=se.section_id,
        exam_id=se.exam_id,
        is_online_active=se.is_online_active,
        is_grades_released=se.is_grades_released,
        is_started=se.is_started,
        release_mode=se.release_mode or "FULL",
        created_at=se.created_at,
        exam=ExamOut.model_validate(se.exam) if se.exam else None,
        submission_count=0,
        average_score=None,
        max_score=exam.total_questions,
    )


@router.delete("/sections/{section_id}/exams/{exam_id}")
async def unassign_exam_from_section(
    section_id: str,
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Unassign an exam from a section. Submissions and student accounts are retained."""
    teacher = await _get_teacher(current_user, db)
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec_res.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(SectionExam).where(
            SectionExam.section_id == section_id,
            SectionExam.exam_id == exam_id,
        )
    )
    se = result.scalar_one_or_none()
    if not se:
        raise HTTPException(status_code=404, detail="Exam is not assigned to this section")

    await db.delete(se)
    await db.commit()
    return {"message": "Exam unassigned from section successfully"}


@router.put("/sections/{section_id}/exams/{exam_id}/toggle-online", response_model=SectionExamOut)
async def toggle_section_exam_online(
    section_id: str,
    exam_id: str,
    body: Optional[SectionExamOnlineToggle] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Toggle online web quiz active status for this specific section."""
    teacher = await _get_teacher(current_user, db)
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec_res.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(SectionExam)
        .options(
            selectinload(SectionExam.exam).selectinload(Exam.answer_keys),
            selectinload(SectionExam.exam).selectinload(Exam.questions),
        )
        .where(
            SectionExam.section_id == section_id,
            SectionExam.exam_id == exam_id,
        )
    )
    se = result.scalar_one_or_none()
    if not se:
        raise HTTPException(status_code=404, detail="Exam is not assigned to this section")

    if body is not None:
        se.is_online_active = body.is_online_active
    else:
        se.is_online_active = not se.is_online_active
    if not se.is_online_active:
        # Closing the lobby also un-releases the exam and clears waiting students.
        # Entries carrying a teacher-granted retake are kept so the attempt counter survives.
        se.is_started = False
        await db.execute(
            delete(ExamLobbyEntry).where(
                ExamLobbyEntry.section_exam_id == se.id,
                ExamLobbyEntry.retakes_granted == 0,
            )
        )

    await db.commit()

    return SectionExamOut(
        id=se.id,
        section_id=se.section_id,
        exam_id=se.exam_id,
        is_online_active=se.is_online_active,
        is_grades_released=se.is_grades_released,
        is_started=se.is_started,
        release_mode=se.release_mode or "FULL",
        created_at=se.created_at,
        exam=ExamOut.model_validate(se.exam) if se.exam else None,
    )


@router.put("/sections/{section_id}/exams/{exam_id}/toggle-grades", response_model=SectionExamOut)
async def toggle_section_exam_grades(
    section_id: str,
    exam_id: str,
    body: Optional[SectionExamGradesToggle] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Toggle releasing grades and questionnaire reviewer for this specific section."""
    teacher = await _get_teacher(current_user, db)
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec_res.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(SectionExam)
        .options(
            selectinload(SectionExam.exam).selectinload(Exam.answer_keys),
            selectinload(SectionExam.exam).selectinload(Exam.questions),
        )
        .where(
            SectionExam.section_id == section_id,
            SectionExam.exam_id == exam_id,
        )
    )
    se = result.scalar_one_or_none()
    if not se:
        raise HTTPException(status_code=404, detail="Exam is not assigned to this section")

    if body is not None and (body.is_grades_released is not None or body.release_mode is not None):
        if body.is_grades_released is not None:
            se.is_grades_released = body.is_grades_released
        if body.release_mode is not None:
            se.release_mode = body.release_mode
    else:
        se.is_grades_released = not se.is_grades_released

    # Update all submissions of students in this section for this exam
    students_res = await db.execute(
        select(Student.id).where(Student.section_id == section_id)
    )
    section_student_ids = list(students_res.scalars().all())

    if section_student_ids:
        await db.execute(
            update(ExamSubmission)
            .where(
                ExamSubmission.exam_id == exam_id,
                ExamSubmission.student_id.in_(section_student_ids),
            )
            .values(is_published=se.is_grades_released)
        )

    await db.commit()

    return SectionExamOut(
        id=se.id,
        section_id=se.section_id,
        exam_id=se.exam_id,
        is_online_active=se.is_online_active,
        is_grades_released=se.is_grades_released,
        is_started=se.is_started,
        release_mode=se.release_mode or "FULL",
        created_at=se.created_at,
        exam=ExamOut.model_validate(se.exam) if se.exam else None,
    )


# ─── Students ─────────────────────────────────────────────────

@router.get("/students", response_model=list[StudentOut])
async def list_all_students(
    search: Optional[str] = Query(None),
    unassigned_only: bool = Query(False),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    List all students registered in the application.
    Supports searching by name, username, or student number.
    Can filter by unassigned_only=True.
    """
    stmt = (
        select(Student)
        .options(joinedload(Student.user), joinedload(Student.section))
    )
    if unassigned_only:
        stmt = stmt.where(Student.section_id.is_(None))
    if search:
        s = f"%{search.strip()}%"
        stmt = stmt.join(User, Student.user_id == User.id).where(
            or_(
                User.first_name.ilike(s),
                User.last_name.ilike(s),
                User.username.ilike(s),
                func.cast(Student.student_number, String).ilike(s),
            )
        )
    stmt = stmt.order_by(Student.student_number)
    result = await db.execute(stmt)
    students = result.scalars().unique().all()

    # Query active pending reset requests
    pending_res = await db.execute(
        select(PasswordResetRequest.student_id).where(
            PasswordResetRequest.status == ResetRequestStatus.PENDING
        )
    )
    pending_ids = set(pending_res.scalars().all())

    out: list[StudentOut] = []
    for s in students:
        s_out = StudentOut.model_validate(s)
        s_out.has_pending_reset_request = s.id in pending_ids
        out.append(s_out)
    return out


@router.post("/students", response_model=StudentOut)
async def create_student(
    body: StudentCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    Register a student into the system.
    Can be created independently without a section roster (section_id=None).
    """
    teacher = await _get_teacher(current_user, db)

    section_id = body.section_id
    if section_id:
        sec = await db.execute(
            select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
        )
        if not sec.scalar_one_or_none():
            raise HTTPException(status_code=404, detail="Section not found")

    # Determine student number (6-digit OMR ID)
    if body.student_number is not None and str(body.student_number).strip():
        try:
            student_number = clean_student_omr_number(body.student_number)
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))
        exists = await db.execute(select(Student).where(Student.student_number == student_number))
        if exists.scalar_one_or_none():
            raise HTTPException(
                status_code=400,
                detail=f"Student ID with OMR digit {student_number:06d} ({STATIC_STUDENT_PREFIX}{student_number:06d}) already exists."
            )
    else:
        student_number = generate_student_number()
        while True:
            exists = await db.execute(select(Student).where(Student.student_number == student_number))
            if not exists.scalar_one_or_none():
                break
            student_number = generate_student_number()

    temp_password = generate_temp_password()
    username = body.username or generate_student_username(body.last_name, student_number)

    # Ensure unique username
    suffix = 0
    orig_username = username
    while True:
        exists = await db.execute(select(User).where(User.username == username))
        if not exists.scalar_one_or_none():
            break
        suffix += 1
        username = f"{orig_username}{suffix}"

    user = User(
        username=username,
        password_hash=hash_password(temp_password),
        role=Role.STUDENT,
        first_name=body.first_name,
        last_name=body.last_name,
        must_change_password=True,
    )
    db.add(user)
    await db.flush()

    student = Student(
        user_id=user.id,
        student_number=student_number,
        section_id=section_id,
    )
    db.add(student)
    await db.flush()

    result = await db.execute(
        select(Student)
        .options(selectinload(Student.user), selectinload(Student.section))
        .where(Student.id == student.id)
    )
    student_out = result.scalar_one()
    out = StudentOut.model_validate(student_out)
    out.temp_password = temp_password
    return out


async def _batch_create_students(
    rows: Sequence[tuple[str, str, Union[int, str, None]]],
    section_id: Optional[str],
    db: AsyncSession,
) -> StudentBulkCreateResponse:
    """
    Shared batch account creator. Each row is (first_name, last_name, student_number|None).
    Rows with invalid/duplicate explicit IDs are skipped and reported in `errors`
    instead of aborting the whole batch.
    """
    created: list[StudentOut] = []
    errors: list[str] = []
    used_numbers: set[int] = set()

    for idx, (first_name, last_name, raw_id) in enumerate(rows, start=1):
        first_name = (first_name or "").strip()
        last_name = (last_name or "").strip()
        label = f"Row {idx} ({first_name} {last_name}".strip() + ")"
        if not first_name or not last_name:
            errors.append(f"Row {idx}: first and last name are required.")
            continue

        student_number: Optional[int] = None
        if raw_id is not None and str(raw_id).strip():
            try:
                student_number = clean_student_omr_number(raw_id)
            except ValueError as e:
                errors.append(f"{label}: {e}")
                continue
            if student_number in used_numbers:
                errors.append(f"{label}: duplicate ID {student_number:06d} in this batch.")
                continue
            exists = await db.execute(select(Student).where(Student.student_number == student_number))
            if exists.scalar_one_or_none():
                errors.append(
                    f"{label}: ID {STATIC_STUDENT_PREFIX}{student_number:06d} is already registered."
                )
                continue
        else:
            student_number = generate_student_number()
            while True:
                exists = await db.execute(select(Student).where(Student.student_number == student_number))
                if not exists.scalar_one_or_none() and student_number not in used_numbers:
                    break
                student_number = generate_student_number()
        used_numbers.add(student_number)

        temp_password = generate_temp_password()
        username = generate_student_username(last_name, student_number)
        suffix = 0
        orig_username = username
        while True:
            exists = await db.execute(select(User).where(User.username == username))
            if not exists.scalar_one_or_none():
                break
            suffix += 1
            username = f"{orig_username}{suffix}"

        user = User(
            username=username,
            password_hash=hash_password(temp_password),
            role=Role.STUDENT,
            first_name=first_name,
            last_name=last_name,
            must_change_password=True,
        )
        db.add(user)
        await db.flush()

        student = Student(user_id=user.id, student_number=student_number, section_id=section_id)
        db.add(student)
        await db.flush()

        result = await db.execute(
            select(Student)
            .options(selectinload(Student.user), selectinload(Student.section))
            .where(Student.id == student.id)
        )
        out = StudentOut.model_validate(result.scalar_one())
        out.temp_password = temp_password
        created.append(out)

    return StudentBulkCreateResponse(
        created=len(created),
        skipped=len(errors),
        total=len(rows),
        students=created,
        errors=errors or None,
    )


async def _verify_optional_section(section_id: Optional[str], current_user: User, db: AsyncSession) -> Optional[str]:
    if not section_id:
        return None
    teacher = await _get_teacher(current_user, db)
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")
    return section_id


@router.post("/students/batch", response_model=StudentBulkCreateResponse)
async def batch_create_students(
    body: StudentBatchCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Batch-register many student accounts at once (optionally enrolling them into a section)."""
    section_id = await _verify_optional_section(body.section_id, current_user, db)
    rows = [(r.first_name, r.last_name, r.student_number) for r in body.students]
    return await _batch_create_students(rows, section_id, db)


@router.post("/students/batch-upload", response_model=StudentBulkCreateResponse)
async def batch_upload_students(
    file: UploadFile = File(...),
    section_id: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Batch-register student accounts from an Excel/CSV file (columns: ID, Name). Section is optional."""
    section_id = await _verify_optional_section(section_id, current_user, db)
    content = await file.read()
    parsed = parse_student_file(content, filename=file.filename or "")
    if not parsed:
        raise HTTPException(
            status_code=400,
            detail="No valid student records found in file. Please ensure the file contains an 'ID' and 'Name' column.",
        )
    rows = [(first, last, csv_id) for first, last, csv_id in parsed]
    return await _batch_create_students(rows, section_id, db)


@router.get("/sections/{section_id}/students", response_model=list[StudentOut])
async def list_students(
    section_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    # Verify section ownership
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(Student)
        .options(selectinload(Student.user), selectinload(Student.section))
        .where(Student.section_id == section_id)
        .order_by(Student.student_number)
    )
    students = result.scalars().all()

    # Query active pending reset requests
    pending_res = await db.execute(
        select(PasswordResetRequest.student_id).where(
            PasswordResetRequest.status == ResetRequestStatus.PENDING
        )
    )
    pending_ids = set(pending_res.scalars().all())

    out: list[StudentOut] = []
    for s in students:
        s_out = StudentOut.model_validate(s)
        s_out.has_pending_reset_request = s.id in pending_ids
        out.append(s_out)
    return out


@router.post("/sections/{section_id}/students", response_model=StudentOut)
async def add_student(
    section_id: str,
    body: StudentCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    # Determine student number (6-digit OMR ID) first so it can form the guaranteed unique username
    if body.student_number is not None and str(body.student_number).strip():
        try:
            student_number = clean_student_omr_number(body.student_number)
        except ValueError as e:
            raise HTTPException(status_code=400, detail=str(e))
        exists = await db.execute(select(Student).where(Student.student_number == student_number))
        if exists.scalar_one_or_none():
            raise HTTPException(
                status_code=400,
                detail=f"Student ID with OMR digit {student_number:06d} ({STATIC_STUDENT_PREFIX}{student_number:06d}) already exists."
            )
    else:
        student_number = generate_student_number()
        while True:
            exists = await db.execute(select(Student).where(Student.student_number == student_number))
            if not exists.scalar_one_or_none():
                break
            student_number = generate_student_number()

    temp_password = generate_temp_password()
    # Student username format: {last_name}.{student_number} (e.g. dilag.282222)
    username = body.username or generate_student_username(body.last_name, student_number)

    # Ensure unique username
    suffix = 0
    orig_username = username
    while True:
        exists = await db.execute(select(User).where(User.username == username))
        if not exists.scalar_one_or_none():
            break
        suffix += 1
        username = f"{orig_username}{suffix}"

    user = User(
        username=username,
        password_hash=hash_password(temp_password),
        role=Role.STUDENT,
        first_name=body.first_name,
        last_name=body.last_name,
        must_change_password=True,
    )
    db.add(user)
    await db.flush()

    student = Student(
        user_id=user.id,
        student_number=student_number,
        section_id=section_id,
    )
    db.add(student)
    await db.flush()

    # Reload with user and section relationships
    result = await db.execute(
        select(Student)
        .options(selectinload(Student.user), selectinload(Student.section))
        .where(Student.id == student.id)
    )
    student_out = result.scalar_one()
    out = StudentOut.model_validate(student_out)
    out.temp_password = temp_password
    return out


@router.post("/sections/{section_id}/students/csv", response_model=StudentBulkCreateResponse)
@router.post("/sections/{section_id}/students/excel", response_model=StudentBulkCreateResponse)
@router.post("/sections/{section_id}/students/upload", response_model=StudentBulkCreateResponse)
async def add_students_csv(
    section_id: str,
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Bulk-add students from an Excel (.xlsx / .xls) or CSV file with columns: ID, Name."""
    teacher = await _get_teacher(current_user, db)
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    content = await file.read()
    rows = parse_student_file(content, filename=file.filename or "")

    if not rows:
        raise HTTPException(
            status_code=400,
            detail="No valid student records found in file. Please ensure the file contains an 'ID' and 'Name' column.",
        )

    created_students = []
    used_numbers = set()
    for first_name, last_name, csv_id in rows:
        student_number = None
        if csv_id is not None and csv_id not in used_numbers:
            exists = await db.execute(select(Student).where(Student.student_number == csv_id))
            if not exists.scalar_one_or_none():
                student_number = csv_id
                used_numbers.add(csv_id)

        if student_number is None:
            student_number = generate_student_number()
            while True:
                exists = await db.execute(select(Student).where(Student.student_number == student_number))
                if not exists.scalar_one_or_none() and student_number not in used_numbers:
                    break
                student_number = generate_student_number()
            used_numbers.add(student_number)

        temp_password = generate_temp_password()
        # Student username format: {last_name}.{student_number} (e.g. dilag.282222)
        username = generate_student_username(last_name, student_number)

        suffix = 0
        orig_username = username
        while True:
            exists = await db.execute(select(User).where(User.username == username))
            if not exists.scalar_one_or_none():
                break
            suffix += 1
            username = f"{orig_username}{suffix}"

        user = User(
            username=username,
            password_hash=hash_password(temp_password),
            role=Role.STUDENT,
            first_name=first_name,
            last_name=last_name,
            must_change_password=True,
        )
        db.add(user)
        await db.flush()

        student = Student(
            user_id=user.id,
            student_number=student_number,
            section_id=section_id,
        )
        db.add(student)
        await db.flush()

        result = await db.execute(
            select(Student)
            .options(selectinload(Student.user), selectinload(Student.section))
            .where(Student.id == student.id)
        )
        student_obj = result.scalar_one()
        out = StudentOut.model_validate(student_obj)
        out.temp_password = temp_password
        created_students.append(out)

    return StudentBulkCreateResponse(created=len(created_students), students=created_students)


@router.post("/sections/{section_id}/enroll", response_model=StudentOut)
async def enroll_student(
    section_id: str,
    body: EnrollStudent,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Enroll an existing student into a section."""
    teacher = await _get_teacher(current_user, db)
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(Student)
        .options(selectinload(Student.user), selectinload(Student.section))
        .where(Student.id == body.student_id)
    )
    student = result.scalar_one_or_none()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    student.section_id = section_id
    await db.flush()

    result = await db.execute(
        select(Student)
        .options(selectinload(Student.user), selectinload(Student.section))
        .where(Student.id == student.id)
    )
    return result.scalar_one()


@router.delete("/sections/{section_id}/students/{student_id}")
async def unenroll_student(
    section_id: str,
    student_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    Remove/unenroll a student from a section roster.
    CRITICAL: Does NOT delete the student account or submissions.
    Only sets student.section_id = None so their account, credentials, and grades remain safe.
    """
    teacher = await _get_teacher(current_user, db)
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(Student).where(Student.id == student_id, Student.section_id == section_id)
    )
    student = result.scalar_one_or_none()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found in this section")

    student.section_id = None
    await db.flush()
    return {"message": "Student unenrolled from section successfully"}


@router.delete("/students/{student_id}")
async def delete_student(
    student_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Permanently delete a student account and user from the system."""
    result = await db.execute(
        select(Student).where(Student.id == student_id)
    )
    student = result.scalar_one_or_none()
    if not student:
        raise HTTPException(status_code=404, detail="Student not found")

    # Delete user (cascades to student)
    user_result = await db.execute(select(User).where(User.id == student.user_id))
    user = user_result.scalar_one_or_none()
    if user:
        await db.delete(user)
    else:
        await db.delete(student)
    await db.flush()
    return {"message": "Student deleted successfully"}


@router.get("/password-reset-requests", response_model=list[PasswordResetRequestOut])
async def list_password_reset_requests(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    List all pending password reset requests for students.
    Includes student info, ID, username, and their enrolled section.
    """
    teacher = await _get_teacher(current_user, db)

    # Get teacher's section IDs
    sec_res = await db.execute(
        select(Section.id).where(Section.teacher_id == teacher.id)
    )
    teacher_sec_ids = set(sec_res.scalars().all())

    stmt = (
        select(PasswordResetRequest)
        .options(
            joinedload(PasswordResetRequest.student).joinedload(Student.user),
            joinedload(PasswordResetRequest.student).joinedload(Student.section),
        )
        .where(PasswordResetRequest.status == ResetRequestStatus.PENDING)
        .order_by(PasswordResetRequest.created_at.desc())
    )
    result = await db.execute(stmt)
    requests = result.scalars().unique().all()

    out: list[PasswordResetRequestOut] = []
    for req in requests:
        st = req.student
        if not st or not st.user:
            continue
        # Show requests if student is in this teacher's section OR if student is unassigned
        if st.section_id is not None and len(teacher_sec_ids) > 0 and st.section_id not in teacher_sec_ids:
            continue

        out.append(PasswordResetRequestOut(
            id=req.id,
            student_id=st.id,
            status=req.status.value,
            created_at=req.created_at,
            resolved_at=req.resolved_at,
            student_name=f"{st.user.first_name} {st.user.last_name}",
            student_username=st.user.username,
            student_number=st.student_number,
            full_student_id=f"02000{st.student_number:06d}",
            section_name=st.section.name if st.section else None,
            section_subject=st.section.subject if st.section else None,
        ))
    return out


@router.post("/students/{student_id}/reset-password", response_model=StudentResetPasswordResponse)
async def reset_student_password(
    student_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    Teacher resets student password:
    - Generates a new temporary password
    - Updates user password_hash
    - Sets must_change_password=True
    - Marks any pending reset requests for this student as RESOLVED
    - Returns the temporary password for the teacher to copy/share with the student
    """
    result = await db.execute(
        select(Student)
        .options(joinedload(Student.user), joinedload(Student.section))
        .where(Student.id == student_id)
    )
    student = result.scalar_one_or_none()
    if not student or not student.user:
        raise HTTPException(status_code=404, detail="Student not found")

    new_temp_password = generate_temp_password()
    student.user.password_hash = hash_password(new_temp_password)
    student.user.must_change_password = True

    # Mark pending requests as resolved
    pending_res = await db.execute(
        select(PasswordResetRequest).where(
            PasswordResetRequest.student_id == student.id,
            PasswordResetRequest.status == ResetRequestStatus.PENDING,
        )
    )
    pending_reqs = pending_res.scalars().all()
    now = datetime.now(timezone.utc)
    for pr in pending_reqs:
        pr.status = ResetRequestStatus.RESOLVED
        pr.resolved_at = now
        pr.resolved_by = current_user.id

    await db.flush()

    return StudentResetPasswordResponse(
        student_id=student.id,
        student_name=f"{student.user.first_name} {student.user.last_name}",
        username=student.user.username,
        full_student_id=f"02000{student.student_number:06d}",
        temp_password=new_temp_password,
        message=f"Password for {student.user.first_name} {student.user.last_name} successfully reset.",
    )


@router.post("/password-reset-requests/{request_id}/dismiss")
async def dismiss_password_reset_request(
    request_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Dismiss a pending password reset request."""
    result = await db.execute(
        select(PasswordResetRequest).where(PasswordResetRequest.id == request_id)
    )
    req = result.scalar_one_or_none()
    if not req:
        raise HTTPException(status_code=404, detail="Reset request not found")

    req.status = ResetRequestStatus.DISMISSED
    req.resolved_at = datetime.now(timezone.utc)
    req.resolved_by = current_user.id
    await db.flush()
    return {"message": "Reset request dismissed successfully"}


# ─── Exams ────────────────────────────────────────────────────

@router.get("/exams", response_model=list[ExamOut])
async def list_exams(
    include_public: bool = Query(False, description="Include public exams from other teachers"),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    options = [selectinload(Exam.answer_keys), selectinload(Exam.questions)]
    if include_public:
        query = select(Exam).options(*options).where(
            (Exam.teacher_id == teacher.id) | (Exam.is_public == True)
        )
    else:
        query = select(Exam).options(*options).where(
            Exam.teacher_id == teacher.id
        )
    result = await db.execute(query.order_by(Exam.created_at.desc()))
    exams = result.scalars().all()
    out = []
    for exam in exams:
        item = ExamOut.model_validate(exam)
        item.is_owner = (exam.teacher_id == teacher.id)
        out.append(item)
    return out


MAX_TEACHER_EXAM_GENERATIONS = 5
TEACHER_QUOTA_PERIOD_DAYS = 21  # 3 weeks (21 days)


async def _get_teacher_quota_status(teacher_id: str, db: AsyncSession) -> dict:
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(days=TEACHER_QUOTA_PERIOD_DAYS)
    stmt = (
        select(TeacherExamGenerationLog)
        .where(
            TeacherExamGenerationLog.teacher_id == teacher_id,
            TeacherExamGenerationLog.created_at >= cutoff,
        )
        .order_by(TeacherExamGenerationLog.created_at.asc())
    )
    res = await db.execute(stmt)
    logs = res.scalars().all()
    used = len(logs)
    remaining = max(0, MAX_TEACHER_EXAM_GENERATIONS - used)

    resets_in_days = None
    if used > 0:
        oldest = logs[0]
        oldest_dt = oldest.created_at if oldest.created_at.tzinfo else oldest.created_at.replace(tzinfo=timezone.utc)
        reset_time = oldest_dt + timedelta(days=TEACHER_QUOTA_PERIOD_DAYS)
        diff_sec = max(0, (reset_time - now).total_seconds())
        resets_in_days = max(1, math.ceil(diff_sec / 86400))

    return {
        "limit": MAX_TEACHER_EXAM_GENERATIONS,
        "used": used,
        "remaining": remaining,
        "period_days": TEACHER_QUOTA_PERIOD_DAYS,
        "period_description": "every 3 weeks",
        "resets_in_days": resets_in_days,
    }


@router.get("/exams/ai-quota")
async def get_teacher_ai_exam_quota(
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Returns teacher's 3-week AI exam generation quota status (5 generations every 3 weeks)."""
    teacher = await _get_teacher(current_user, db)
    return await _get_teacher_quota_status(teacher.id, db)


class GenerateExamFromTextRequest(BaseModel):
    raw_text: str
    question_count: int = 10
    title: Optional[str] = None
    include_scenarios: bool = True
    custom_instructions: Optional[str] = None


@router.post("/exams/generate-from-handout-text")
async def generate_exam_from_handout_text(
    body: GenerateExamFromTextRequest,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    Teacher AI Exam Generator:
    Generates structured test questions, 4 options (A-D), answer keys, and explanations
    directly from lecture notes or syllabus text using Google Gemini AI.
    Enforces quota of 5 generations every 3 weeks.
    """
    teacher = await _get_teacher(current_user, db)
    quota = await _get_teacher_quota_status(teacher.id, db)
    if quota["remaining"] <= 0:
        days_str = f"{quota['resets_in_days']} day" if quota.get("resets_in_days") == 1 else f"{quota.get('resets_in_days', 21)} days"
        raise HTTPException(
            status_code=429,
            detail=f"The AI Exam Generator is currently unavailable and taking a rest for this cycle. It will be ready for you again in {days_str}. In the meantime, you can freely create exams manually or edit existing questions anytime without any limits!"
        )

    if not body.raw_text or len(body.raw_text.strip()) < 30:
        raise HTTPException(
            status_code=400,
            detail="Please provide at least a few sentences or paragraphs of lesson notes."
        )

    count = max(3, min(50, body.question_count))
    from app.services.gemini_service import generate_exam_from_handout
    result = generate_exam_from_handout(
        body.raw_text,
        question_count=count,
        custom_title=body.title,
        include_scenarios=body.include_scenarios,
        custom_instructions=body.custom_instructions,
    )

    if result and result.get("questions"):
        log_entry = TeacherExamGenerationLog(
            teacher_id=teacher.id,
            title=result.get("title") or "AI Exam",
            question_count=len(result["questions"]),
        )
        db.add(log_entry)
        await db.commit()

        updated_quota = await _get_teacher_quota_status(teacher.id, db)
        result["quota_used"] = updated_quota["used"]
        result["quota_remaining"] = updated_quota["remaining"]
        result["quota_limit"] = MAX_TEACHER_EXAM_GENERATIONS
        result["quota_resets_in_days"] = updated_quota["resets_in_days"]

    return result


@router.post("/exams/generate-from-handout-file")
async def generate_exam_from_handout_file(
    file: UploadFile = File(...),
    question_count: int = Form(10),
    title: Optional[str] = Form(None),
    include_scenarios: bool = Form(True),
    custom_instructions: Optional[str] = Form(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """
    Teacher AI Exam Generator:
    Uploads a PDF, TXT, or Markdown handout and automatically formulates exam questions,
    answer keys, and explanations using Google Gemini AI.
    Enforces quota of 5 generations every 3 weeks.
    """
    teacher = await _get_teacher(current_user, db)
    quota = await _get_teacher_quota_status(teacher.id, db)
    if quota["remaining"] <= 0:
        days_str = f"{quota['resets_in_days']} day" if quota.get("resets_in_days") == 1 else f"{quota.get('resets_in_days', 21)} days"
        raise HTTPException(
            status_code=429,
            detail=f"The AI Exam Generator is currently unavailable and taking a rest for this cycle. It will be ready for you again in {days_str}. In the meantime, you can freely create exams manually or edit existing questions anytime without any limits!"
        )

    filename = (file.filename or "").lower()
    if not (filename.endswith(".pdf") or filename.endswith(".txt") or filename.endswith(".docx") or filename.endswith(".md")):
        raise HTTPException(
            status_code=400,
            detail="Please upload a PDF (.pdf), Text file (.txt), or Markdown (.md)."
        )

    content_bytes = await file.read()
    if not content_bytes:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    extracted_text = ""
    if filename.endswith(".pdf"):
        from app.reviewer_generator import extract_text_from_pdf
        try:
            extracted_text = extract_text_from_pdf(content_bytes)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not read PDF: {str(e)}")
    else:
        try:
            extracted_text = content_bytes.decode("utf-8", errors="replace")
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"Could not decode file: {str(e)}")

    if not extracted_text or len(extracted_text.strip()) < 30:
        raise HTTPException(
            status_code=400,
            detail="The document does not contain enough readable text to formulate questions."
        )

    clean_file_title = title or (file.filename or "Handout Exam").rsplit(".", 1)[0].replace("_", " ").replace("-", " ").title()
    count = max(3, min(50, question_count))
    from app.services.gemini_service import generate_exam_from_handout
    result = generate_exam_from_handout(
        extracted_text,
        question_count=count,
        custom_title=clean_file_title,
        include_scenarios=include_scenarios,
        custom_instructions=custom_instructions,
    )

    if result and result.get("questions"):
        log_entry = TeacherExamGenerationLog(
            teacher_id=teacher.id,
            title=result.get("title") or clean_file_title,
            question_count=len(result["questions"]),
        )
        db.add(log_entry)
        await db.commit()

        updated_quota = await _get_teacher_quota_status(teacher.id, db)
        result["quota_used"] = updated_quota["used"]
        result["quota_remaining"] = updated_quota["remaining"]
        result["quota_limit"] = MAX_TEACHER_EXAM_GENERATIONS
        result["quota_resets_in_days"] = updated_quota["resets_in_days"]

    return result


@router.post("/exams", response_model=ExamOut)
async def create_exam(
    body: ExamCreate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)

    exam = Exam(
        teacher_id=teacher.id,
        title=body.title,
        total_questions=body.total_questions,
        is_public=body.is_public or body.is_grades_released,
        is_online_active=body.is_online_active,
        is_grades_released=body.is_grades_released,
        shuffle_questions=body.shuffle_questions,
        shuffle_options=body.shuffle_options,
        has_time_limit=body.has_time_limit,
        time_limit_minutes=body.time_limit_minutes,
    )
    db.add(exam)
    await db.flush()

    for ak in body.answer_keys:
        key = ExamAnswerKey(
            exam_id=exam.id,
            question_number=ak.question_number,
            correct_option=ak.correct_option,
        )
        db.add(key)

    if body.questions:
        for q in body.questions:
            db.add(ExamQuestion(
                exam_id=exam.id,
                question_number=q.question_number,
                question_text=q.question_text,
                option_a=q.option_a,
                option_b=q.option_b,
                option_c=q.option_c,
                option_d=q.option_d,
                explanation=q.explanation,
            ))
    await db.flush()

    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(Exam.id == exam.id)
    )
    return result.scalar_one()


@router.get("/exams/{exam_id}", response_model=ExamOut)
async def get_exam(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(
            Exam.id == exam_id,
            (Exam.teacher_id == teacher.id) | (Exam.is_public == True)
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
    return exam


@router.put("/exams/{exam_id}", response_model=ExamOut)
async def update_exam(
    exam_id: str,
    body: ExamUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    if body.title is not None:
        exam.title = body.title
    if body.total_questions is not None:
        exam.total_questions = body.total_questions
    if body.is_online_active is not None:
        exam.is_online_active = body.is_online_active
    if body.shuffle_questions is not None:
        exam.shuffle_questions = body.shuffle_questions
    if body.shuffle_options is not None:
        exam.shuffle_options = body.shuffle_options
    if body.has_time_limit is not None:
        exam.has_time_limit = body.has_time_limit
    if body.time_limit_minutes is not None:
        exam.time_limit_minutes = body.time_limit_minutes
    if body.is_grades_released is not None:
        exam.is_grades_released = body.is_grades_released
        exam.is_public = body.is_grades_released
        subs_res = await db.execute(
            select(ExamSubmission).where(ExamSubmission.exam_id == exam.id)
        )
        for s in subs_res.scalars().all():
            s.is_published = body.is_grades_released
    elif body.is_public is not None:
        exam.is_public = body.is_public
        exam.is_grades_released = body.is_public
        subs_res = await db.execute(
            select(ExamSubmission).where(ExamSubmission.exam_id == exam.id)
        )
        for s in subs_res.scalars().all():
            s.is_published = body.is_public

    if body.answer_keys is not None:
        for key in exam.answer_keys:
            await db.delete(key)
        await db.flush()

        for ak in body.answer_keys:
            db.add(ExamAnswerKey(
                exam_id=exam.id,
                question_number=ak.question_number,
                correct_option=ak.correct_option,
            ))
        await db.flush()

    if body.questions is not None:
        for q in exam.questions:
            await db.delete(q)
        await db.flush()

        for q in body.questions:
            db.add(ExamQuestion(
                exam_id=exam.id,
                question_number=q.question_number,
                question_text=q.question_text,
                option_a=q.option_a,
                option_b=q.option_b,
                option_c=q.option_c,
                option_d=q.option_d,
                explanation=q.explanation,
            ))
        await db.flush()

    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(Exam.id == exam.id)
    )
    return result.scalar_one()


@router.put("/exams/{exam_id}/toggle-online", response_model=ExamOut)
@router.patch("/exams/{exam_id}/online", response_model=ExamOut)
async def toggle_exam_online(
    exam_id: str,
    body: Optional[ExamOnlineToggle] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Toggle whether students can take this exam online via web."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    if body is not None:
        exam.is_online_active = body.is_online_active
    else:
        exam.is_online_active = not exam.is_online_active

    await db.commit()
    return exam


@router.put("/exams/{exam_id}/toggle-grades-release", response_model=ExamOut)
@router.patch("/exams/{exam_id}/grades-release", response_model=ExamOut)
async def toggle_exam_grades_release(
    exam_id: str,
    body: Optional[ExamGradesReleaseToggle] = None,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Toggle releasing exam grades, shaded answers, and questionnaire reviewer to students."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    if body is not None:
        exam.is_grades_released = body.is_grades_released
    else:
        exam.is_grades_released = not exam.is_grades_released

    exam.is_public = exam.is_grades_released

    # Update all submissions: set is_published to match grade release status
    await db.execute(
        update(ExamSubmission)
        .where(ExamSubmission.exam_id == exam.id)
        .values(is_published=exam.is_grades_released)
    )

    await db.commit()
    return exam


@router.patch("/exams/{exam_id}/visibility", response_model=ExamOut)
async def toggle_exam_visibility(
    exam_id: str,
    body: ExamVisibilityUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Toggle an exam between Public (visible to students with reviewer) and Private."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys), selectinload(Exam.questions)).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    exam.is_public = body.is_public
    exam.is_grades_released = body.is_public
    await db.execute(
        update(ExamSubmission)
        .where(ExamSubmission.exam_id == exam.id)
        .values(is_published=body.is_public)
    )

    await db.commit()
    return exam


@router.put("/exams/{exam_id}/questionnaire", response_model=list[QuestionOut])
async def update_questionnaire(
    exam_id: str,
    body: QuestionnaireUpdate,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Save or update the questionnaire (questions, choices, explanations) for an exam separately from the answer keys."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(selectinload(Exam.questions), selectinload(Exam.answer_keys)).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    for q in exam.questions:
        await db.delete(q)
    await db.flush()

    for q in body.questions:
        db.add(ExamQuestion(
            exam_id=exam.id,
            question_number=q.question_number,
            question_text=q.question_text,
            option_a=q.option_a,
            option_b=q.option_b,
            option_c=q.option_c,
            option_d=q.option_d,
            explanation=q.explanation,
        ))
    await db.flush()

    if body.answer_keys is not None:
        for key in exam.answer_keys:
            await db.delete(key)
        await db.flush()

        for ak in body.answer_keys:
            db.add(ExamAnswerKey(
                exam_id=exam.id,
                question_number=ak.question_number,
                correct_option=ak.correct_option,
            ))
        await db.flush()

    await db.commit()

    result = await db.execute(
        select(ExamQuestion).where(ExamQuestion.exam_id == exam.id).order_by(ExamQuestion.question_number)
    )
    return result.scalars().all()


@router.get("/exams/{exam_id}/questionnaire", response_model=list[QuestionOut])
async def get_questionnaire(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Retrieve questionnaire for an exam."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).where(
            Exam.id == exam_id,
            (Exam.teacher_id == teacher.id) | (Exam.is_public == True)
        )
    )
    if not result.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Exam not found")

    result = await db.execute(
        select(ExamQuestion).where(ExamQuestion.exam_id == exam_id).order_by(ExamQuestion.question_number)
    )
    return result.scalars().all()


@router.delete("/exams/{exam_id}")
async def delete_exam(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).where(Exam.id == exam_id, Exam.teacher_id == teacher.id)
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")
    await db.delete(exam)
    await db.flush()
    return {"message": "Exam deleted successfully"}


@router.post("/exams/{exam_id}/clone", response_model=ExamOut)
async def clone_exam(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Clone a public exam (including answer keys and questionnaire) into the current teacher's exams."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).options(
            selectinload(Exam.answer_keys),
            selectinload(Exam.questions)
        ).where(
            Exam.id == exam_id, Exam.is_public == True
        )
    )
    source = result.scalar_one_or_none()
    if not source:
        raise HTTPException(status_code=404, detail="Public exam not found")

    new_exam = Exam(
        teacher_id=teacher.id,
        title=f"{source.title} (Copy)",
        total_questions=source.total_questions,
        is_public=False,
        shuffle_questions=source.shuffle_questions,
        shuffle_options=source.shuffle_options,
        has_time_limit=source.has_time_limit,
        time_limit_minutes=source.time_limit_minutes,
    )
    db.add(new_exam)
    await db.flush()

    for key in source.answer_keys:
        db.add(ExamAnswerKey(
            exam_id=new_exam.id,
            question_number=key.question_number,
            correct_option=key.correct_option,
        ))

    for q in source.questions:
        db.add(ExamQuestion(
            exam_id=new_exam.id,
            question_number=q.question_number,
            question_text=q.question_text,
            option_a=q.option_a,
            option_b=q.option_b,
            option_c=q.option_c,
            option_d=q.option_d,
            explanation=q.explanation,
        ))
    await db.flush()

    result = await db.execute(
        select(Exam).options(
            selectinload(Exam.answer_keys),
            selectinload(Exam.questions)
        ).where(Exam.id == new_exam.id)
    )
    return result.scalar_one()


# ─── PDF Generation ───────────────────────────────────────────

@router.get("/exams/{exam_id}/bubble-sheet")
async def generate_bubble_sheet(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Generate a printable PDF bubble sheet for the exam."""
    from fastapi.responses import StreamingResponse
    from app.services.pdf_generator import generate_bubble_sheet_pdf

    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(Exam).where(Exam.id == exam_id, Exam.teacher_id == teacher.id)
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    pdf_bytes = generate_bubble_sheet_pdf(exam.title, exam.total_questions)

    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="bubble_sheet_{exam_id}.pdf"'},
    )


@router.get("/blank-bubble-sheet")
async def generate_blank_bubble_sheet(
    questions: int = Query(50, description="Total questions"),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Generate a generic printable PDF bubble sheet."""
    from fastapi.responses import StreamingResponse
    from app.services.pdf_generator import generate_bubble_sheet_pdf

    pdf_bytes = generate_bubble_sheet_pdf("Blank Answer Sheet", questions)

    return StreamingResponse(
        iter([pdf_bytes]),
        media_type="application/pdf",
        headers={"Content-Disposition": 'attachment; filename="blank_bubble_sheet.pdf"'},
    )


# ─── OMR Scanning ─────────────────────────────────────────────

@router.post("/scan-key")
async def scan_answer_key(
    file: UploadFile = File(...),
    total_questions: int = Query(50),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Upload a scanned bubble sheet image to extract answers without saving to DB."""
    from app.omr.scanner import process_omr_sheet

    image_bytes = await file.read()
    try:
        omr_result = process_omr_sheet(image_bytes, total_questions)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"OMR processing failed: {str(e)}")

    # Return the dictionary of detected answers
    return {"answers": omr_result["answers"]}


@router.post("/exams/{exam_id}/scan", response_model=SubmissionOut)
async def scan_sheet(
    exam_id: str,
    file: UploadFile = File(...),
    section_id: Optional[str] = Query(None),
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Upload a scanned bubble sheet image and process it through the OMR pipeline."""
    from app.omr.scanner import process_omr_sheet
    from app.services.gdrive_service import upload_to_gdrive

    teacher = await _get_teacher(current_user, db)

    # Verify exam ownership
    result = await db.execute(
        select(Exam).options(selectinload(Exam.answer_keys)).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    exam = result.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    if not exam.answer_keys:
        raise HTTPException(status_code=400, detail="Exam has no answer keys defined")

    # If section_id is provided, verify section ownership
    target_section = None
    if section_id:
        sec_res = await db.execute(
            select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
        )
        target_section = sec_res.scalar_one_or_none()
        if not target_section:
            raise HTTPException(status_code=404, detail="Selected section not found")

    # Read image
    image_bytes = await file.read()

    # Run OMR pipeline
    try:
        omr_result = process_omr_sheet(image_bytes, exam.total_questions)
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"OMR processing failed: {str(e)}")

    # Find student by detected student number
    student_number = omr_result["student_number"]
    result = await db.execute(
        select(Student)
        .options(selectinload(Student.user), selectinload(Student.section))
        .where(
            Student.student_number == student_number
        )
    )
    student = result.scalar_one_or_none()
    if not student:
        raise HTTPException(
            status_code=404,
            detail=f"No student found with student number {student_number}"
        )

    # If a specific section was targeted, ensure student belongs to that section
    if target_section and student.section_id != target_section.id:
        current_sec_name = student.section.name if student.section else "No Section"
        st_name = f"{student.user.first_name} {student.user.last_name}" if student.user else f"Student #{student.student_number}"
        raise HTTPException(
            status_code=400,
            detail=f"Student {st_name} ({student.full_student_id}) belongs to section '{current_sec_name}', not '{target_section.name}'."
        )

    # Upload image to Google Drive
    try:
        image_url = upload_to_gdrive(image_bytes, f"scan_{exam_id}_{student_number}.jpg")
    except Exception:
        image_url = None  # Continue without image URL if GDrive fails

    # Build answer key lookup
    key_map = {ak.question_number: ak.correct_option for ak in exam.answer_keys}

    # Grade and calculate score
    total_score = 0
    answers = []
    for q_num, detected in omr_result["answers"].items():
        q_int = int(q_num)
        correct = key_map.get(q_int, "")
        is_correct = detected == correct if detected else False
        if is_correct:
            total_score += 1
        answers.append({
            "question_number": q_int,
            "shaded_option": detected,
            "is_correct": is_correct,
        })

    # Check for existing submission to upsert cleanly
    existing_sub_res = await db.execute(
        select(ExamSubmission).where(
            ExamSubmission.exam_id == exam.id,
            ExamSubmission.student_id == student.id
        )
    )
    existing_sub = existing_sub_res.scalar_one_or_none()

    if existing_sub:
        submission = existing_sub
        submission.total_score = total_score
        submission.max_score = exam.total_questions
        if image_url:
            submission.scanned_image_url = image_url
        submission.scanned_at = datetime.now(timezone.utc)
        # Clear previous answers
        await db.execute(
            delete(SubmissionAnswer).where(SubmissionAnswer.submission_id == submission.id)
        )
    else:
        submission = ExamSubmission(
            exam_id=exam.id,
            student_id=student.id,
            total_score=total_score,
            max_score=exam.total_questions,
            scanned_image_url=image_url,
            is_published=exam.is_grades_released,
        )
        db.add(submission)
        await db.flush()

    for ans in answers:
        db.add(SubmissionAnswer(
            submission_id=submission.id,
            question_number=ans["question_number"],
            shaded_option=ans["shaded_option"],
            is_correct=ans["is_correct"],
        ))
    await db.flush()

    # Reload with relationships
    result = await db.execute(
        select(ExamSubmission)
        .options(
            selectinload(ExamSubmission.answers),
            selectinload(ExamSubmission.student).selectinload(Student.user),
            selectinload(ExamSubmission.student).selectinload(Student.section),
            selectinload(ExamSubmission.exam).selectinload(Exam.answer_keys),
        )
        .where(ExamSubmission.id == submission.id)
    )
    return result.scalar_one()


# ─── Section Exam Gradebook & Excel Export ───────────────────

@router.get("/sections/{section_id}/exams/{exam_id}/gradebook")
async def get_section_exam_gradebook(
    section_id: str,
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Retrieve full section roster with grades or ABSENT status for a given exam."""
    teacher = await _get_teacher(current_user, db)

    # Verify section
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    section = sec_res.scalar_one_or_none()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    # Verify exam
    exam_res = await db.execute(
        select(Exam).where(Exam.id == exam_id, Exam.teacher_id == teacher.id)
    )
    exam = exam_res.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    # Fetch all students in this section
    students_res = await db.execute(
        select(Student)
        .options(selectinload(Student.user))
        .where(Student.section_id == section_id)
        .order_by(Student.student_number.asc())
    )
    students = students_res.scalars().all()

    # Fetch all submissions for this exam from students in this section
    student_ids = [s.id for s in students]
    submissions_map = {}
    if student_ids:
        subs_res = await db.execute(
            select(ExamSubmission)
            .options(selectinload(ExamSubmission.answers))
            .where(
                ExamSubmission.exam_id == exam_id,
                ExamSubmission.student_id.in_(student_ids)
            )
        )
        for sub in subs_res.scalars().all():
            submissions_map[sub.student_id] = sub

    records = []
    total_enrolled = len(students)
    graded_count = 0
    total_scores = 0

    for st in students:
        sub = submissions_map.get(st.id)
        u = st.user
        student_name = f"{u.last_name}, {u.first_name}" if u else f"Student #{st.student_number}"
        full_student_id = st.full_student_id

        if sub:
            graded_count += 1
            total_scores += sub.total_score
            pct = round((sub.total_score / sub.max_score * 100), 1) if sub.max_score > 0 else 0.0
            status = "GRADED"
            score = sub.total_score
            max_s = sub.max_score
            scanned_at = sub.scanned_at.isoformat() if sub.scanned_at else None
            submission_id = sub.id
            scanned_image_url = sub.scanned_image_url
            is_published = sub.is_published
        else:
            status = "ABSENT"
            score = 0
            max_s = exam.total_questions
            pct = 0.0
            scanned_at = None
            submission_id = None
            scanned_image_url = None
            is_published = False

        records.append({
            "student_id": st.id,
            "student_number": st.student_number,
            "full_student_id": full_student_id,
            "student_name": student_name,
            "status": status,
            "score": score,
            "max_score": max_s,
            "percentage": pct,
            "scanned_at": scanned_at,
            "submission_id": submission_id,
            "scanned_image_url": scanned_image_url,
            "is_published": is_published,
        })

    absent_count = total_enrolled - graded_count
    class_average = round(total_scores / graded_count, 1) if graded_count > 0 else 0.0
    average_percentage = round((class_average / exam.total_questions * 100), 1) if exam.total_questions > 0 else 0.0

    return {
        "section_id": section.id,
        "section_name": section.name,
        "subject": section.subject,
        "exam_id": exam.id,
        "exam_title": exam.title,
        "total_questions": exam.total_questions,
        "total_enrolled": total_enrolled,
        "graded_count": graded_count,
        "absent_count": absent_count,
        "class_average": class_average,
        "average_percentage": average_percentage,
        "records": records,
    }


@router.get("/sections/{section_id}/exams/{exam_id}/export-excel")
async def export_section_exam_excel(
    section_id: str,
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Download styled Excel (.xlsx) file containing complete section roster, scores, and absent records."""
    import io
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    teacher = await _get_teacher(current_user, db)

    # Verify section
    sec_res = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    section = sec_res.scalar_one_or_none()
    if not section:
        raise HTTPException(status_code=404, detail="Section not found")

    # Verify exam
    exam_res = await db.execute(
        select(Exam).where(Exam.id == exam_id, Exam.teacher_id == teacher.id)
    )
    exam = exam_res.scalar_one_or_none()
    if not exam:
        raise HTTPException(status_code=404, detail="Exam not found")

    # Fetch students & submissions
    students_res = await db.execute(
        select(Student)
        .options(selectinload(Student.user))
        .where(Student.section_id == section_id)
        .order_by(Student.student_number.asc())
    )
    students = students_res.scalars().all()

    student_ids = [s.id for s in students]
    submissions_map = {}
    if student_ids:
        subs_res = await db.execute(
            select(ExamSubmission)
            .where(
                ExamSubmission.exam_id == exam_id,
                ExamSubmission.student_id.in_(student_ids)
            )
        )
        for sub in subs_res.scalars().all():
            submissions_map[sub.student_id] = sub

    from typing import Any
    wb = openpyxl.Workbook()
    ws: Any = wb.active
    ws.title = f"Gradebook - {section.name}"[:31]

    # Colors & Styles
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="0F172A")
    subtitle_font = Font(name="Calibri", size=11, bold=False, color="475569")
    meta_label_font = Font(name="Calibri", size=10, bold=True, color="334155")
    meta_val_font = Font(name="Calibri", size=10, color="0F172A")

    absent_fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    absent_font = Font(name="Calibri", size=10, bold=True, color="DC2626")
    graded_font = Font(name="Calibri", size=10, bold=True, color="16A34A")

    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title & Metadata
    ws["A1"] = "EVALO - EXAM GRADEBOOK & ATTENDANCE RECORD"
    ws["A1"].font = title_font
    ws["A2"] = f"Subject: {section.subject} | Section: {section.name} | Exam: {exam.title}"
    ws["A2"].font = subtitle_font

    ws["A4"] = "Section:"
    ws["A4"].font = meta_label_font
    ws["B4"] = section.name
    ws["B4"].font = meta_val_font

    ws["C4"] = "Exam Title:"
    ws["C4"].font = meta_label_font
    ws["D4"] = exam.title
    ws["D4"].font = meta_val_font

    ws["E4"] = "Total Questions:"
    ws["E4"].font = meta_label_font
    ws["F4"] = exam.total_questions
    ws["F4"].font = meta_val_font

    ws["A5"] = "Total Enrolled:"
    ws["A5"].font = meta_label_font
    ws["B5"] = len(students)
    ws["B5"].font = meta_val_font

    graded_count = sum(1 for s in students if s.id in submissions_map)
    absent_count = len(students) - graded_count

    ws["C5"] = "Graded / Present:"
    ws["C5"].font = meta_label_font
    ws["D5"] = graded_count
    ws["D5"].font = meta_val_font

    ws["E5"] = "Absent:"
    ws["E5"].font = meta_label_font
    ws["F5"] = absent_count
    ws["F5"].font = meta_val_font

    # Table Headers
    headers = [
        "No.", "Student ID", "Full Name", "Status",
        "Score", "Max Score", "Percentage", "Remarks", "Scanned / Submitted At"
    ]
    start_row = 7
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    # Data Rows
    current_row = start_row + 1
    for idx, st in enumerate(students, 1):
        sub = submissions_map.get(st.id)
        u = st.user
        full_name = f"{u.last_name}, {u.first_name}" if u else f"Student #{st.student_number}"

        if sub:
            status = "GRADED"
            score = sub.total_score
            max_s = sub.max_score
            pct = round((score / max_s * 100), 1) if max_s > 0 else 0.0
            pct_str = f"{pct}%"
            remarks = "PASSED" if pct >= 50 else "NEEDS IMPROVEMENT"
            sub_date = sub.scanned_at.strftime("%Y-%m-%d %H:%M") if sub.scanned_at else "Scanned"
        else:
            status = "ABSENT"
            score = 0
            max_s = exam.total_questions
            pct_str = "0.0%"
            remarks = "ABSENT"
            sub_date = "-"

        row_vals = [
            idx,
            st.full_student_id,
            full_name,
            status,
            score,
            max_s,
            pct_str,
            remarks,
            sub_date
        ]

        for col_idx, val in enumerate(row_vals, 1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.border = thin_border
            if col_idx in [1, 2, 4, 5, 6, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

            if status == "ABSENT":
                cell.fill = absent_fill
                if col_idx == 4 or col_idx == 8:
                    cell.font = absent_font
            else:
                if col_idx == 4:
                    cell.font = graded_font

        current_row += 1

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            if cell.row < 7:
                continue
            if len(val) > max_len:
                max_len = len(val)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws.column_dimensions["A"].width = 8
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["I"].width = 22

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    clean_sec = "".join(c for c in section.name if c.isalnum() or c in (' ', '_', '-')).strip()
    clean_exam = "".join(c for c in exam.title if c.isalnum() or c in (' ', '_', '-')).strip()
    filename = f"Gradebook_{clean_sec}_{clean_exam}.xlsx"

    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'}
    )


# ─── Submissions ──────────────────────────────────────────────

@router.get("/exams/{exam_id}/submissions", response_model=list[SubmissionOut])
async def list_submissions(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(ExamSubmission)
        .options(
            selectinload(ExamSubmission.answers),
            selectinload(ExamSubmission.student).selectinload(Student.user),
            selectinload(ExamSubmission.student).selectinload(Student.section),
        )
        .join(Exam)
        .where(Exam.id == exam_id, Exam.teacher_id == teacher.id)
        .order_by(ExamSubmission.scanned_at.desc())
    )
    return result.scalars().all()


@router.put("/submissions/{submission_id}/override", response_model=SubmissionOut)
async def override_submission(
    submission_id: str,
    body: SubmissionOverride,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Manually override detected answers for a submission."""
    teacher = await _get_teacher(current_user, db)

    result = await db.execute(
        select(ExamSubmission)
        .options(
            selectinload(ExamSubmission.answers),
            selectinload(ExamSubmission.exam).selectinload(Exam.answer_keys),
        )
        .join(Exam)
        .where(ExamSubmission.id == submission_id, Exam.teacher_id == teacher.id)
    )
    submission = result.scalar_one_or_none()
    if not submission or not submission.exam:
        raise HTTPException(status_code=404, detail="Submission or exam not found")

    key_map = {ak.question_number: ak.correct_option for ak in submission.exam.answer_keys}

    # Delete old answers
    for ans in submission.answers:
        await db.delete(ans)
    await db.flush()

    total_score = 0
    for item in body.answers:
        correct = key_map.get(item.question_number, "")
        is_correct = item.correct_option == correct
        if is_correct:
            total_score += 1
        db.add(SubmissionAnswer(
            submission_id=submission.id,
            question_number=item.question_number,
            shaded_option=item.correct_option,
            is_correct=is_correct,
        ))

    submission.total_score = total_score
    await db.flush()

    result = await db.execute(
        select(ExamSubmission)
        .options(
            selectinload(ExamSubmission.answers),
            selectinload(ExamSubmission.student).selectinload(Student.user),
            selectinload(ExamSubmission.student).selectinload(Student.section),
            selectinload(ExamSubmission.exam).selectinload(Exam.answer_keys),
        )
        .where(ExamSubmission.id == submission.id)
    )
    return result.scalar_one()


@router.put("/submissions/{submission_id}/publish")
async def publish_submission(
    submission_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(ExamSubmission).join(Exam).where(
            ExamSubmission.id == submission_id, Exam.teacher_id == teacher.id
        )
    )
    submission = result.scalar_one_or_none()
    if not submission:
        raise HTTPException(status_code=404, detail="Submission not found")

    submission.is_published = True
    await db.flush()
    return {"message": "Grades published successfully"}


@router.put("/exams/{exam_id}/publish-all")
async def publish_all_submissions(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Publish all submissions for an exam."""
    teacher = await _get_teacher(current_user, db)
    result = await db.execute(
        select(ExamSubmission).join(Exam).where(
            Exam.id == exam_id, Exam.teacher_id == teacher.id
        )
    )
    submissions = result.scalars().all()
    for s in submissions:
        s.is_published = True
    await db.flush()
    return {"message": f"Published {len(submissions)} submissions"}


# ─── Gradebook / Analytics ────────────────────────────────────

@router.get("/exams/{exam_id}/analytics", response_model=list[ItemAnalysis])
async def get_item_analysis(
    exam_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Get per-question item analysis for an exam."""
    teacher = await _get_teacher(current_user, db)

    result = await db.execute(
        select(ExamSubmission)
        .options(selectinload(ExamSubmission.answers))
        .join(Exam)
        .where(Exam.id == exam_id, Exam.teacher_id == teacher.id)
    )
    submissions = result.scalars().all()

    if not submissions:
        return []

    # Aggregate by question
    question_data = {}
    for sub in submissions:
        for ans in sub.answers:
            qn = ans.question_number
            if qn not in question_data:
                question_data[qn] = {"correct": 0, "total": 0, "options": {}}
            question_data[qn]["total"] += 1
            if ans.is_correct:
                question_data[qn]["correct"] += 1
            opt = ans.shaded_option or "BLANK"
            question_data[qn]["options"][opt] = question_data[qn]["options"].get(opt, 0) + 1

    analysis = []
    for qn in sorted(question_data.keys()):
        d = question_data[qn]
        analysis.append(ItemAnalysis(
            question_number=qn,
            correct_count=d["correct"],
            total_count=d["total"],
            correct_percentage=round(d["correct"] / d["total"] * 100, 1) if d["total"] else 0,
            option_distribution=d["options"],
        ))
    return analysis


@router.get("/sections/{section_id}/gradebook")
async def get_section_gradebook(
    section_id: str,
    db: AsyncSession = Depends(get_db),
    current_user: User = Depends(require_role(Role.TEACHER)),
):
    """Get grade summary for all students in a section."""
    teacher = await _get_teacher(current_user, db)
    sec = await db.execute(
        select(Section).where(Section.id == section_id, Section.teacher_id == teacher.id)
    )
    if not sec.scalar_one_or_none():
        raise HTTPException(status_code=404, detail="Section not found")

    result = await db.execute(
        select(Student)
        .options(
            selectinload(Student.user),
            selectinload(Student.submissions).selectinload(ExamSubmission.exam),
        )
        .where(Student.section_id == section_id)
        .order_by(Student.student_number)
    )
    students = result.scalars().all()

    gradebook = []
    for student in students:
        if not student.user:
            continue
        published_subs = [
            s for s in student.submissions
            if (s.is_published or (s.exam and s.exam.is_public)) and s.exam
        ]
        avg_score = (
            sum(s.total_score for s in published_subs) / len(published_subs)
            if published_subs else 0
        )
        gradebook.append({
            "student_id": student.id,
            "student_number": student.student_number,
            "name": f"{student.user.first_name} {student.user.last_name}",
            "total_exams_taken": len(published_subs),
            "average_score": round(avg_score, 1),
            "exams": [
                {
                    "exam_title": s.exam.title,
                    "score": s.total_score,
                    "max_score": s.max_score,
                    "percentage": round(s.total_score / s.max_score * 100, 1) if s.max_score else 0,
                }
                for s in published_subs
                if s.exam
            ],
        })
    return gradebook
